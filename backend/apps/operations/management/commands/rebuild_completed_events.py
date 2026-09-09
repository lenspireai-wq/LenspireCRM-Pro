from datetime import date, datetime

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from apps.operations.models import CalendarEvent
from apps.operations.views import normalize_excel_time


IMPORT_FIELDS = (
    "title", "client_name", "event_type", "start_date", "start_time", "city", "status",
    "handled_by", "couple_name", "contact_no", "photo", "video", "candid", "cinematic",
    "drone", "assistant", "bts", "notes", "date_status", "tbd_month",
)


class Command(BaseCommand):
    help = "Archive active completed events and rebuild them exactly from a workbook."

    def add_arguments(self, parser):
        parser.add_argument("--workbook", required=True, help="Path to the completed-events workbook.")
        parser.add_argument("--expected-count", type=int, required=True)
        parser.add_argument("--apply", action="store_true")

    @staticmethod
    def text(value):
        return "" if value is None else str(value).strip()

    def load_rows(self, workbook_path):
        sheet = load_workbook(workbook_path, data_only=True, read_only=True).active
        iterator = sheet.iter_rows(values_only=True)
        headers = [self.text(value) for value in next(iterator)]
        rows = []
        for row_number, values in enumerate(iterator, start=2):
            row = dict(zip(headers, values))
            if not row.get("title"):
                continue
            payload = {field: row.get(field) for field in IMPORT_FIELDS if field in row}
            for field in CalendarEvent._meta.fields:
                if field.name not in payload or field.get_internal_type() not in {"CharField", "TextField"}:
                    continue
                payload[field.name] = self.text(payload[field.name])
            start_date = payload.get("start_date")
            if isinstance(start_date, datetime):
                payload["start_date"] = start_date.date()
            elif start_date is not None and not isinstance(start_date, date):
                raise CommandError(f"Row {row_number}: start_date is not a valid date.")
            if "start_time" in payload:
                normalized_time = normalize_excel_time(payload["start_time"])
                if normalized_time is None:
                    time_note = self.text(payload.pop("start_time"))
                    notes = self.text(payload.get("notes"))
                    payload["notes"] = f"{notes}\nTime details: {time_note}".strip()
                else:
                    payload["start_time"] = normalized_time or None
            payload["status"] = "Completed"
            payload["is_archived"] = False
            rows.append(payload)
        return rows

    def handle(self, *args, **options):
        rows = self.load_rows(options["workbook"])
        expected_count = options["expected_count"]
        if len(rows) != expected_count:
            raise CommandError(f"Workbook has {len(rows)} usable rows, not {expected_count}.")

        completed = CalendarEvent.objects.filter(status__iexact="Completed", is_archived=False)
        organizations = list(completed.values_list("organization_id", flat=True).distinct())
        if len(organizations) != 1:
            raise CommandError("Expected completed events in exactly one active organization. No records were changed.")
        organization_id = organizations[0]
        current_count = completed.filter(organization_id=organization_id).count()
        self.stdout.write(
            f"organization={organization_id} active_completed={current_count} workbook_rows={len(rows)} "
            f"archive={current_count} create={len(rows)}"
        )
        if not options["apply"]:
            self.stdout.write(self.style.WARNING("Dry run only; no records were changed."))
            return

        with transaction.atomic():
            completed.filter(organization_id=organization_id).update(is_archived=True)
            CalendarEvent.objects.bulk_create(
                [CalendarEvent(organization_id=organization_id, **row) for row in rows],
                batch_size=250,
            )
            rebuilt_count = CalendarEvent.objects.filter(
                organization_id=organization_id, status__iexact="Completed", is_archived=False,
            ).count()
            if rebuilt_count != expected_count:
                raise CommandError(
                    f"Rebuild created {rebuilt_count} completed events instead of {expected_count}; transaction rolled back."
                )
        self.stdout.write(self.style.SUCCESS(
            f"Archived {current_count} previous completed events and rebuilt {expected_count} active completed events."
        ))
