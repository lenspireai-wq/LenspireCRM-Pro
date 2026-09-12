from datetime import date, time, timedelta

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook
from rest_framework.test import APIClient

from apps.core.models import Organization
from apps.sales.models import Booking, Customer, Lead
from apps.users.models import User
from .models import CalendarEvent, PhotographerDetail


class OperationsApiTests(TestCase):
    def setUp(self):
        self.organization = Organization.objects.create(name="Studio", slug="studio")
        self.user = User.objects.create_user(username="admin", password="secret", organization=self.organization)
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def test_event_and_photographer_are_organization_scoped(self):
        response = self.client.post("/api/events/", {"title": "Wedding", "client_name": "Asha", "event_type": "Wedding", "start_date": "2026-09-20", "date_status": "Confirmed"})
        self.assertEqual(response.status_code, 201)
        self.assertEqual(CalendarEvent.objects.get().organization, self.organization)
        response = self.client.post("/api/photographers/", {"name": "Avi", "status": "Available"})
        self.assertEqual(response.status_code, 201)
        self.assertEqual(PhotographerDetail.objects.get().organization, self.organization)

    def confirmed_booking(self, *, name="Asha Patel", couple_name="Asha & Rohan", phone="9876543210"):
        lead = Lead.objects.create(
            organization=self.organization,
            lead_code=f"LD-{name.replace(' ', '-').upper()}",
            name=name,
            couple_name=couple_name,
            mobile=phone,
            event_type="Wedding",
            status="Confirmed",
        )
        customer = Customer.objects.create(
            organization=self.organization,
            customer_code=f"C-{lead.id}",
            name=name,
            phone=phone,
            lead=lead,
        )
        return Booking.objects.create(
            organization=self.organization,
            booking_code=f"B-{lead.id}",
            customer=customer,
            lead=lead,
            event_type="Wedding",
            status="Confirmed",
            quoted_amount="100000.00",
        )

    def test_manual_event_auto_links_to_one_confirmed_booking(self):
        booking = self.confirmed_booking()

        response = self.client.post(
            "/api/events/",
            {
                "title": "Asha Patel · Engagement",
                "client_name": "Asha Patel",
                "couple_name": "Asha & Rohan",
                "contact_no": "98765 43210",
                "event_type": "Engagement",
                "start_date": "2030-01-10",
                "date_status": "Confirmed",
            },
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["booking"], booking.id)
        self.assertEqual(response.data["customer"], booking.customer_id)

    def test_batch_link_preview_is_read_only_until_applied(self):
        booking = self.confirmed_booking()
        event = CalendarEvent.objects.create(
            organization=self.organization,
            title="Asha Patel · Engagement",
            client_name="Asha Patel",
            couple_name="Asha & Rohan",
            contact_no="9876543210",
            event_type="Engagement",
            start_date=date(2030, 1, 10),
        )

        preview = self.client.get("/api/events/link-confirmed-bookings/")
        self.assertEqual(preview.status_code, 200, preview.data)
        self.assertTrue(preview.data["dry_run"])
        self.assertEqual(preview.data["summary"]["eligible"], 1)
        event.refresh_from_db()
        self.assertIsNone(event.booking_id)

        applied = self.client.post("/api/events/link-confirmed-bookings/", {"apply": True}, format="json")
        self.assertEqual(applied.status_code, 200, applied.data)
        self.assertEqual(applied.data["summary"]["linked"], 1)
        event.refresh_from_db()
        self.assertEqual(event.booking_id, booking.id)
        self.assertEqual(event.customer_id, booking.customer_id)

    def test_manual_event_stays_unlinked_when_confirmed_match_is_ambiguous(self):
        self.confirmed_booking(name="Asha Patel", couple_name="Asha & Rohan", phone="9876543210")
        self.confirmed_booking(name="Asha Patel Two", couple_name="Asha & Rohan", phone="9123456789")

        response = self.client.post(
            "/api/events/",
            {
                "title": "Asha · Engagement",
                "couple_name": "Asha & Rohan",
                "event_type": "Engagement",
                "start_date": "2030-01-10",
                "date_status": "Confirmed",
            },
        )

        self.assertEqual(response.status_code, 201, response.data)
        self.assertIsNone(response.data["booking"])

    def test_photographers_stay_scoped_for_a_superuser(self):
        other_organization = Organization.objects.create(name="Other Studio", slug="other-studio")
        PhotographerDetail.objects.create(organization=self.organization, name="Avi")
        PhotographerDetail.objects.create(organization=other_organization, name="Outside crew")
        self.user.is_superuser = True
        self.user.is_staff = True
        self.user.save(update_fields=("is_superuser", "is_staff"))

        response = self.client.get("/api/photographers/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data["results"]), 1)
        self.assertEqual(response.data["results"][0]["name"], "Avi")

    def test_tbd_event_requires_month(self):
        response = self.client.post("/api/events/", {"title": "Wedding", "event_type": "Wedding", "date_status": "TBD Month"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("tbd_month", response.data)

    def test_past_event_is_automatically_completed(self):
        yesterday = date.today() - timedelta(days=1)
        response = self.client.post(
            "/api/events/",
            {
                "title": "Past Wedding",
                "event_type": "Wedding",
                "start_date": yesterday.isoformat(),
                "date_status": "Confirmed",
                "status": "Scheduled",
            },
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["status"], "Completed")

    def test_future_event_cannot_remain_completed(self):
        tomorrow = date.today() + timedelta(days=1)
        event = CalendarEvent.objects.create(
            organization=self.organization,
            title="Future Wedding",
            event_type="Wedding",
            start_date=tomorrow,
            status="Completed",
        )
        response = self.client.get("/api/events/")
        self.assertEqual(response.status_code, 200)
        event.refresh_from_db()
        self.assertEqual(event.status, "Scheduled")

    def test_cancelled_past_event_stays_cancelled(self):
        yesterday = date.today() - timedelta(days=1)
        event = CalendarEvent.objects.create(
            organization=self.organization,
            title="Cancelled Wedding",
            event_type="Wedding",
            start_date=yesterday,
            status="Cancelled",
        )
        self.client.get("/api/events/")
        event.refresh_from_db()
        self.assertEqual(event.status, "Cancelled")

    def test_event_today_is_automatically_in_progress(self):
        response = self.client.post(
            "/api/events/",
            {
                "title": "Wedding Today",
                "event_type": "Wedding",
                "start_date": timezone.localdate().isoformat(),
                "date_status": "Confirmed",
            },
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["status"], "In Progress")

    def test_ready_future_event_is_automatically_confirmed(self):
        tomorrow = timezone.localdate() + timedelta(days=1)
        response = self.client.post(
            "/api/events/",
            {
                "title": "Ready Wedding",
                "event_type": "Wedding",
                "start_date": tomorrow.isoformat(),
                "start_time": "10:00:00",
                "date_status": "Confirmed",
                "city": "Mumbai Venue",
                "notes": "Arrive thirty minutes early.",
                "photo": "Aakash Gorde · 918421258470",
                "video": "NA",
                "candid": "NA",
                "cinematic": "NA",
                "drone": "NA",
                "assistant": "NA",
                "bts": "NA",
            },
        )
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["status"], "Confirmed")

    def test_rejects_duplicate_events_and_photographers(self):
        event = {"title": "Wedding", "client_name": "Asha", "event_type": "Wedding", "start_date": "2026-09-20", "date_status": "Confirmed"}
        self.assertEqual(self.client.post("/api/events/", event).status_code, 201)
        self.assertEqual(self.client.post("/api/events/", event).status_code, 400)
        self.assertEqual(self.client.post("/api/photographers/", {"name": "Avi", "mobile": "+91 98765 43210"}).status_code, 201)
        self.assertEqual(self.client.post("/api/photographers/", {"name": "Avi Duplicate", "mobile": "9876543210"}).status_code, 400)

    def test_upcoming_events_excel_export(self):
        CalendarEvent.objects.create(organization=self.organization, title="Wedding", start_date=date(2026, 9, 20))
        response = self.client.get("/api/events/export/")
        self.assertEqual(response.status_code, 200)
        self.assertIn("upcoming-events.xlsx", response["Content-Disposition"])

    def test_upcoming_events_excel_import(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["title", "event_type", "start_date", "date_status"])
        sheet.append(["Reception", "Reception", date(2026, 10, 2), "Confirmed"])
        from io import BytesIO
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("events.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response = self.client.post("/api/events/import/", {"file": upload}, format="multipart")
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["created"], 1)

    def test_upcoming_events_import_accepts_tbd_in_date_column(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["title", "event_type", "start_date", "date_status"])
        sheet.append(["Date pending", "Wedding", "TBD", "Confirmed"])
        from io import BytesIO
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("events.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self.client.post("/api/events/import/", {"file": upload}, format="multipart")

        self.assertEqual(response.status_code, 201, response.data)
        event = CalendarEvent.objects.get(organization=self.organization, title="Date pending")
        self.assertIsNone(event.start_date)
        self.assertEqual(event.date_status, "TBD")

    def test_upcoming_events_import_converts_tbd_month_date_text(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["title", "event_type", "start_date", "date_status"])
        sheet.append(["December pending", "Wedding", "TBD - December", "Confirmed"])
        from io import BytesIO
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("events.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self.client.post("/api/events/import/", {"file": upload}, format="multipart")

        self.assertEqual(response.status_code, 201, response.data)
        event = CalendarEvent.objects.get(organization=self.organization, title="December pending")
        self.assertIsNone(event.start_date)
        self.assertEqual(event.date_status, "TBD Month")
        self.assertEqual(event.tbd_month, "2026-12")

    def test_importing_an_edited_export_updates_the_existing_event(self):
        event = CalendarEvent.objects.create(
            organization=self.organization,
            title="Original wedding",
            client_name="Asha",
            event_type="Wedding",
            start_date=date(2026, 10, 2),
            status="Completed",
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["id", "title", "client_name", "event_type", "start_date", "date_status", "notes"])
        sheet.append([event.id, "Edited wedding", "Asha", "Wedding", date(2026, 10, 2), "Confirmed", "Updated from Excel"])
        from io import BytesIO
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("edited-events.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self.client.post("/api/events/import/", {"file": upload}, format="multipart")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data, {"created": 0, "updated": 1})
        self.assertEqual(CalendarEvent.objects.filter(organization=self.organization).count(), 1)
        event.refresh_from_db()
        self.assertEqual(event.title, "Edited wedding")
        self.assertEqual(event.notes, "Updated from Excel")

    def test_completed_events_import_allows_blank_city_and_notes(self):
        event = CalendarEvent.objects.create(
            organization=self.organization,
            title="Completed wedding",
            client_name="Asha",
            event_type="Wedding",
            start_date=date(2025, 10, 2),
            status="Completed",
            city="Mumbai",
            notes="Previous note",
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["id", "title", "client_name", "event_type", "start_date", "date_status", "city", "notes"])
        sheet.append([event.id, "Completed wedding", "Asha", "Wedding", date(2025, 10, 2), "Confirmed", None, None])
        from io import BytesIO
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("completed-events.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self.client.post("/api/events/import/", {"file": upload}, format="multipart")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data, {"created": 0, "updated": 1})
        event.refresh_from_db()
        self.assertEqual(event.city, "")
        self.assertEqual(event.notes, "")

    def test_completed_events_import_moves_free_text_time_to_notes(self):
        event = CalendarEvent.objects.create(
            organization=self.organization,
            title="Completed wedding",
            client_name="Asha",
            event_type="Wedding",
            start_date=date(2025, 10, 2),
            start_time="10:00:00",
            status="Completed",
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["title", "client_name", "event_type", "start_date", "start_time", "date_status", "notes"])
        sheet.append(["Completed wedding", "Asha", "Wedding", date(2025, 10, 2), "TICKET DONE", "Confirmed", ""])
        from io import BytesIO
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("completed-events.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self.client.post("/api/events/import/", {"file": upload}, format="multipart")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data, {"created": 0, "updated": 1})
        event.refresh_from_db()
        self.assertEqual(event.start_time.strftime("%H:%M:%S"), "10:00:00")
        self.assertEqual(event.notes, "Time details: TICKET DONE")

    def test_completed_events_import_converts_non_string_city_cells(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["title", "event_type", "start_date", "date_status", "city"])
        sheet.append(["Timed city", "Wedding", date(2025, 10, 2), "Confirmed", time(15, 0)])
        from io import BytesIO
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("completed-events.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self.client.post("/api/events/import/", {"file": upload}, format="multipart")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(CalendarEvent.objects.get(title="Timed city").city, "15:00:00")

    def test_importing_a_legacy_export_without_id_updates_a_unique_match(self):
        event = CalendarEvent.objects.create(
            organization=self.organization,
            title="Original wedding",
            client_name="Asha",
            contact_no="9876543210",
            event_type="Wedding",
            start_date=date(2026, 10, 2),
            start_time="10:00:00",
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["title", "client_name", "contact_no", "event_type", "start_date", "start_time", "date_status", "notes"])
        sheet.append(["Edited wedding", "Asha", "9876543210", "Wedding", date(2026, 10, 2), "10:00:00", "Confirmed", "Updated legacy export"])
        from io import BytesIO
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("legacy-events.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self.client.post("/api/events/import/", {"file": upload}, format="multipart")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data, {"created": 0, "updated": 1})
        self.assertEqual(CalendarEvent.objects.filter(organization=self.organization).count(), 1)
        event.refresh_from_db()
        self.assertEqual(event.title, "Edited wedding")
        self.assertEqual(event.notes, "Updated legacy export")

    def test_import_by_id_updates_an_event_when_another_duplicate_exists(self):
        event = CalendarEvent.objects.create(
            organization=self.organization,
            title="Repeated event",
            client_name="Asha",
            contact_no="9876543210",
            event_type="Wedding",
            start_date=date(2026, 10, 2),
        )
        CalendarEvent.objects.create(
            organization=self.organization,
            title="Repeated event",
            client_name="Asha",
            contact_no="9876543210",
            event_type="Wedding",
            start_date=date(2026, 10, 2),
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["id", "title", "client_name", "contact_no", "event_type", "start_date", "date_status", "notes"])
        sheet.append([event.id, "Repeated event", "Asha", "9876543210", "Wedding", date(2026, 10, 2), "Confirmed", "Updated duplicate"])
        from io import BytesIO
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("legacy-events.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self.client.post("/api/events/import/", {"file": upload}, format="multipart")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data, {"created": 0, "updated": 1})
        event.refresh_from_db()
        self.assertEqual(event.notes, "Updated duplicate")

    def test_legacy_import_uses_title_date_and_time_when_phone_format_differs(self):
        event = CalendarEvent.objects.create(
            organization=self.organization,
            title="Ankit Gupta · Pre-wedding",
            client_name="Ankit Gupta",
            contact_no="7977518696",
            event_type="Pre-wedding",
            start_date=date(2026, 9, 8),
            start_time="07:00:00",
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["title", "client_name", "contact_no", "event_type", "start_date", "start_time", "date_status", "notes"])
        sheet.append(["Ankit Gupta · Pre-wedding", "Ankit Gupta", 7977518696, "Pre-wedding", date(2026, 9, 8), "07:00:00", "Confirmed", "Imported note"])
        from io import BytesIO
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("completed-events.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self.client.post("/api/events/import/", {"file": upload}, format="multipart")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data, {"created": 0, "updated": 1})
        event.refresh_from_db()
        self.assertEqual(event.notes, "Imported note")

    def test_legacy_import_updates_all_identical_duplicate_matches(self):
        event_values = {
            "organization": self.organization,
            "title": "Mamta Jadhav · Pre-wedding",
            "client_name": "Mamta Jadhav",
            "contact_no": "8291631219",
            "event_type": "Pre-wedding",
            "start_date": date(2025, 10, 19),
            "start_time": "07:10:00",
        }
        first = CalendarEvent.objects.create(**event_values)
        second = CalendarEvent.objects.create(**event_values)
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["title", "client_name", "contact_no", "event_type", "start_date", "start_time", "date_status", "notes"])
        sheet.append(["Mamta Jadhav · Pre-wedding", "Mamta Jadhav", "8291631219", "Pre-wedding", date(2025, 10, 19), "07:10:00", "Confirmed", "South Bombay"])
        from io import BytesIO
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("completed-events.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

        response = self.client.post("/api/events/import/", {"file": upload}, format="multipart")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data, {"created": 0, "updated": 2})
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertEqual(first.notes, "South Bombay")
        self.assertEqual(second.notes, "South Bombay")
